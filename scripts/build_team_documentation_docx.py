"""
Generate a clear, approachable, and easy-to-read team summary document for Turio (.docx).
Applies the karpathy-clear-explanations principles:
- Familiar words, short active sentences, crystal clear analogies.
- Answers the 'What, Why, How, and Where' for every team member.
- Beautiful visual layout, clean tables, intuitive callouts, and embedded diagrams.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Palette: Tech Modern Friendly
# Primary: Deep Navy Teal (#0E7490)
# Secondary: Cool Slate (#334155)
# Accent: Vibrant Amber (#D97706)
# Body Text: Dark Carbon (#1E293B)
# Subtext: Muted Grey (#64748B)
# Card/Table BG: Light Soft Fog (#F8FAFC)
# Highlight Green: #059669
# Border: #CBD5E1

COLOR_PRIMARY = RGBColor(14, 116, 144)    # #0E7490
COLOR_SECONDARY = RGBColor(51, 65, 85)     # #334155
COLOR_ACCENT = RGBColor(217, 119, 6)       # #D97706
COLOR_TEXT_DARK = RGBColor(30, 41, 59)     # #1E293B
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # #64748B
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_PRIMARY = "0E7490"
HEX_SECONDARY = "334155"
HEX_ACCENT = "D97706"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "F0F9FF"
HEX_WARN_BG = "FEF2F2"
HEX_WARN_BORDER = "EF4444"
HEX_SUCCESS_BG = "ECFDF5"
HEX_SUCCESS_BORDER = "10B981"


def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for edge, border in borders.items():
        if border:
            b = OxmlElement(f'w:{edge}')
            b.set(qn('w:val'), border.get('val', 'single'))
            b.set(qn('w:sz'), str(border.get('sz', 4)))
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), border.get('color', HEX_BORDER))
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{edge}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)


def style_table(table, col_widths, headers, data, alt_shading=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_row = table.rows[0]
    hdr_tr = hdr_row._tr.get_or_add_trPr()
    hdr_tr.append(OxmlElement('w:tblHeader'))
    hdr_tr.append(OxmlElement('w:cantSplit'))

    for i, cell in enumerate(hdr_row.cells):
        cell.width = col_widths[i]
        set_cell_background(cell, HEX_PRIMARY)
        set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_borders(cell,
                         top={'val': 'single', 'sz': 6, 'color': HEX_PRIMARY},
                         bottom={'val': 'single', 'sz': 10, 'color': HEX_ACCENT})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(headers[i])
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_WHITE

    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        row_tr = row._tr.get_or_add_trPr()
        row_tr.append(OxmlElement('w:cantSplit'))
        bg_color = HEX_LIGHT_BG if (alt_shading and r_idx % 2 == 1) else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            if bg_color != "FFFFFF":
                set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell,
                             top={'val': 'single', 'sz': 4, 'color': HEX_BORDER},
                             bottom={'val': 'single', 'sz': 4, 'color': HEX_BORDER})
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(cell_value))
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_TEXT_DARK


def add_callout(doc, text, title=None, style_type="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.5)

    if style_type == "warning":
        bg_hex = HEX_WARN_BG
        border_hex = HEX_WARN_BORDER
        text_color = RGBColor(185, 28, 28)
    elif style_type == "success":
        bg_hex = HEX_SUCCESS_BG
        border_hex = HEX_SUCCESS_BORDER
        text_color = RGBColor(5, 150, 105)
    else: # info
        bg_hex = HEX_CALLOUT_BG
        border_hex = HEX_PRIMARY
        text_color = COLOR_PRIMARY

    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_borders(cell, left={'val': 'single', 'sz': 24, 'color': border_hex})

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    if title:
        r_title = p.add_run(title + "\n")
        r_title.bold = True
        r_title.font.name = 'Calibri'
        r_title.font.size = Pt(10.5)
        r_title.font.color.rgb = text_color

    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_TEXT_DARK

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(6)


def add_page_number_to_run(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    r = run._r
    r.append(fldChar1)
    r.append(instrText)
    r.append(fldChar2)
    r.append(fldChar3)


def add_heading_1(doc, text):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return h


def add_heading_2(doc, text):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return h


def add_heading_3(doc, text):
    h = doc.add_heading(level=3)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.bold = True
    run.font.color.rgb = COLOR_ACCENT
    return h


def add_body_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_TEXT_DARK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_TEXT_DARK
    return p


def add_bullet_point(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = COLOR_TEXT_DARK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_TEXT_DARK
    return p


def build_team_documentation():
    print("Building Team-Friendly Documentation...")
    doc = docx.Document()

    # 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.different_first_page_header_footer = True

        # Header (page 2+)
        header_p = s.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = header_p.add_run("TURIO — Team Project Overview & Architecture Guide")
        r_hdr.font.name = 'Calibri'
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = COLOR_TEXT_MUTED

        # Footer
        footer_p = s.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_ftr_l = footer_p.add_run("College 4-2 Final Project — Team Guide\t\tPage ")
        r_ftr_l.font.name = 'Calibri'
        r_ftr_l.font.size = Pt(8.5)
        r_ftr_l.font.color.rgb = COLOR_TEXT_MUTED
        add_page_number_to_run(footer_p.add_run())

    # ==========================================
    # COVER / HEADER
    # ==========================================
    p_tag = doc.add_paragraph()
    p_tag.paragraph_format.space_before = Pt(30)
    p_tag.paragraph_format.space_after = Pt(6)
    r_tag = p_tag.add_run("PROJECT HANDBOOK & ARCHITECTURE SUMMARY")
    r_tag.bold = True
    r_tag.font.name = 'Calibri'
    r_tag.font.size = Pt(10.5)
    r_tag.font.color.rgb = COLOR_ACCENT

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("TURIO: What It Is, How It Works & How We Built It")
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run(
        "A clear, practical guide for our team members to understand the complete codebase, "
        "machine learning matching engine, database, and job application automation."
    )
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Quick Metadata Card
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Project Goal", "Turn a resume into matched tech jobs and automate applying with 1 click."),
        ("Key Tech Stack", "Python (FastAPI), React / Next.js 14, PostgreSQL 16, SBERT AI, Playwright."),
        ("Current State", "Fully working locally on native Windows. 56,339 jobs imported, 385 tests passing."),
        ("Who This Guide Is For", "Anyone on our project team who needs to present, demo, or code on Turio.")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, HEX_LIGHT_BG)
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_borders(c0, bottom={'val': 'single', 'sz': 4, 'color': HEX_BORDER})
        set_cell_borders(c1, bottom={'val': 'single', 'sz': 4, 'color': HEX_BORDER})

        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.name = 'Calibri'
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_TEXT_DARK

    doc.add_page_break()

    # ==========================================
    # SECTION 1: THE SIMPLE EXPLANATION
    # ==========================================
    add_heading_1(doc, "1. What is Turio in Plain English?")

    add_body_paragraph(
        doc,
        "Applying for tech jobs today is tedious and repetitive. You upload your resume to a job board, "
        "re-type your skills and contact details dozens of times, write repetitive answers to common questions, "
        "and manually track which companies have replied."
    )
    add_body_paragraph(
        doc,
        "Turio solves this by creating one smooth, connected experience from your resume to your submitted application:"
    )

    add_callout(
        doc,
        "1. You upload your resume (PDF).\n"
        "2. Turio reads your skills, projects, and work history automatically (without leaking your data).\n"
        "3. Turio scores 56,000+ real tech jobs using AI to show you your best matches with explanations.\n"
        "4. You click 'Apply', and Turio fills the company's real application form for you.\n"
        "5. You review the filled form on screen, make any edits, and approve it.\n"
        "6. Turio submits the application safely and tracks your progress in your dashboard.",
        title="THE 6-STEP TURIO WORKFLOW",
        style_type="info"
    )

    add_body_paragraph(
        doc,
        "The key philosophy behind Turio is: Always keep the human in the loop. The AI and automation do the boring typing, "
        "but the user always has the final approval before anything is submitted to an employer."
    )

    # Insert System Workflow Architecture Flowchart
    flowchart_path = r"D:\College\4-2 project\docs\images\workflow_architecture.png"
    if os.path.exists(flowchart_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(flowchart_path, width=Inches(6.2))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Figure 1: Visual overview of the 4 phases in the Turio pipeline.")
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SECTION 2: HOW THE SYSTEM IS BUILT (ARCHITECTURE)
    # ==========================================
    doc.add_page_break()
    add_heading_1(doc, "2. How the System is Built (Architecture)")

    add_body_paragraph(
        doc,
        "When explaining this project in an interview or project viva, think of Turio as having 4 main parts that talk to each other:"
    )

    arch_headers = ["Part", "What It Does", "Tech Used", "Where It Lives in Code"]
    arch_widths = [Inches(1.2), Inches(2.2), Inches(1.4), Inches(1.7)]
    arch_data = [
        ["1. Frontend", "The modern web UI where the user uploads resumes, browses jobs, and reviews applications.", "Next.js 14, React, Tailwind CSS", "frontend/ (App Router on Port 13000)"],
        ["2. Backend API", "The brain of the system. Manages authentication, runs matching algorithms, and orchestrates applications.", "Python 3.11, FastAPI, Pydantic", "backend/app/ (FastAPI on Port 18080)"],
        ["3. Database & Storage", "Stores user profiles, jobs catalog, matches, and application history. Private storage holds resume PDFs.", "PostgreSQL 16 & S3rver", "Database Port 18543, S3 Storage Port 19000"],
        ["4. Automation Worker", "Controls a headless browser (Playwright) that visits real employer application pages and fills in answers safely.", "Playwright (Chromium) & Python", "backend/app/adapters/ashby.py & workers/"]
    ]
    tbl_arch = doc.add_table(rows=len(arch_data) + 1, cols=4)
    style_table(tbl_arch, arch_widths, arch_headers, arch_data)

    add_body_paragraph(doc, "", space_after=8)
    add_heading_2(doc, "2.1 Why did we choose this setup?")
    why_points = [
        ("Fast & Lightweight: ", "FastAPI gives us sub-second API responses and automatic OpenAPI documentation (/docs)."),
        ("Native Windows Friendly: ", "Everything runs cleanly on Windows without Docker or WSL headaches using our start-local-stack.ps1 script."),
        ("Secure by Default: ", "User session tokens are stored in HTTP-Only cookies so malicious scripts cannot steal passwords or login tokens.")
    ]
    for b_title, b_desc in why_points:
        add_bullet_point(doc, b_desc, bold_prefix=b_title)

    # ==========================================
    # SECTION 3: THE SMART FEATURES
    # ==========================================
    doc.add_page_break()
    add_heading_1(doc, "3. Deep Dive into the Smart Features")

    add_heading_2(doc, "3.1 Feature A: Smart Resume Parser (No External Leaks)")
    add_body_paragraph(
        doc,
        "Many resumes use tricky formatting: 2 columns, weird spacing in headers like 'E X P E R I E N C E', "
        "and customized section titles. Instead of paying an external AI API to read personal data, "
        "we built our own parser in backend/app/services/resume_layout.py."
    )
    add_bullet_point(doc, "It scans the page layout to separate left and right columns cleanly.")
    add_bullet_point(doc, "It extracts email, phone, LinkedIn, GitHub, skills, degree, and work history.")
    add_bullet_point(doc, "It saves these into your Profile in the database so you only upload once.")

    add_heading_2(doc, "3.2 Feature B: How AI Matches Jobs to You (The Math Made Easy)")
    add_body_paragraph(
        doc,
        "We have 56,339 real tech jobs in our database from 3,448 companies (imported from our dataset). "
        "How do we know which job fits you best? We don't just do basic keyword search. We use a hybrid formula:"
    )

    add_callout(
        doc,
        "Final Match Score = (50% AI Semantic Meaning) + (30% Exact Skills Overlap) + (15% Experience Years) + (5% Education)\n\n"
        "• AI Meaning (50%): Uses Sentence-BERT (all-MiniLM-L6-v2) to turn your resume and the job description into 384 numbers (vectors) and compares how close their meanings are.\n"
        "• Skills Overlap (30%): Checks if you have the required tools (e.g. Python, React, PostgreSQL).\n"
        "• Experience & Education (20%): Compares your years in tech and degree level.",
        title="HOW MATCH SCORES ARE CALCULATED",
        style_type="info"
    )

    add_heading_2(doc, "3.3 Feature C: Automated Form Filling & The 'Answer Once' Superpower")
    add_body_paragraph(
        doc,
        "When you want to apply to a job (for example on an Ashby-hosted career page):"
    )
    fill_steps = [
        ("Form Inspection: ", "Turio opens the real application page in the background and finds every question."),
        ("Auto-Fill from Profile: ", "It puts your name, email, phone, LinkedIn, and original PDF resume in place."),
        ("Answer Once: ", "If an employer asks a question you haven't answered yet (e.g., 'Do you require sponsorship?'), Turio asks you ONCE. When you answer, Turio saves it in saved_answers and uses it for ALL future applications!"),
        ("Review & Approve: ", "You see the filled application preview on your screen. You can edit any answer."),
        ("Guarded Submit: ", "When you click 'Approve & Submit', Turio sends the data and confirms the employer server replied with FormSubmitSuccess.")
    ]
    for b_title, b_desc in fill_steps:
        add_bullet_point(doc, b_desc, bold_prefix=b_title)

    # Insert FSM Diagram
    fsm_path = r"D:\College\4-2 project\docs\images\fsm_lifecycle.png"
    if os.path.exists(fsm_path):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_before = Pt(12)
        p_img2.paragraph_format.space_after = Pt(2)
        run_img2 = p_img2.add_run()
        run_img2.add_picture(fsm_path, width=Inches(6.2))

        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_before = Pt(0)
        p_cap2.paragraph_format.space_after = Pt(12)
        r_cap2 = p_cap2.add_run("Figure 2: The state machine that guides every application safely from draft to submission.")
        r_cap2.font.name = 'Calibri'
        r_cap2.font.size = Pt(8.5)
        r_cap2.font.italic = True
        r_cap2.font.color.rgb = COLOR_TEXT_MUTED

    # ==========================================
    # SECTION 4: THE DATABASE EXPLAINED SIMPLY
    # ==========================================
    doc.add_page_break()
    add_heading_1(doc, "4. The Database Made Simple")

    add_body_paragraph(
        doc,
        "We use PostgreSQL. Instead of memorizing 16 tables, here is how the data connects in 5 simple groups:"
    )

    db_headers = ["Group", "Tables Inside", "What They Store", "Why It Matters"]
    db_widths = [Inches(1.2), Inches(1.8), Inches(2.0), Inches(1.5)]
    db_data = [
        ["1. User & Auth", "users, refresh_tokens", "Account logins, hashed passwords, and secure session tokens.", "Lets users securely sign in and out."],
        ["2. Candidate Profile", "profiles, skills, experiences, educations, resumes", "All your parsed resume details: phone, skills, degree, jobs.", "The central database for filling applications."],
        ["3. Jobs Catalog", "job_sources, jobs", "56,339 tech jobs from Ashby and other platforms.", "Provides the database of opportunities to search and filter."],
        ["4. AI Matches", "job_matches, embeddings", "Calculated match scores (0-100) and stored SBERT vector numbers.", "Powers the 'Recommended for You' tab."],
        ["5. Applications & Tracking", "applications, application_fields, application_answers, saved_answers, application_events", "Draft applications, form questions, your confirmed answers, and submission receipts.", "Powers the application tracker and 'Answer Once' memory."]
    ]
    tbl_db = doc.add_table(rows=len(db_data) + 1, cols=4)
    style_table(tbl_db, db_widths, db_headers, db_data)

    # ==========================================
    # SECTION 5: HOW TO RUN AND DEMO THE PROJECT
    # ==========================================
    doc.add_page_break()
    add_heading_1(doc, "5. How to Run, Test, and Demo Turio")

    add_body_paragraph(
        doc,
        "Here are the exact commands you need to know when running the project or presenting it to evaluators:"
    )

    add_heading_2(doc, "5.1 Starting the Entire Application (One Command)")
    add_callout(
        doc,
        "Open PowerShell as Administrator and run:\n\n"
        "cd \"D:\\College\\4-2 project\"\n"
        "powershell -ExecutionPolicy Bypass -File scripts/start-local-stack.ps1 -SkipBuild\n\n"
        "That's it! This script automatically starts:\n"
        "• PostgreSQL Database (Port 18543)\n"
        "• S3rver Object Storage (Port 19000)\n"
        "• FastAPI Backend (Port 18080)\n"
        "• Next.js Web Frontend (Port 13000)\n"
        "• Background Task Worker",
        title="EASY ONE-COMMAND LAUNCHER",
        style_type="success"
    )

    add_heading_2(doc, "5.2 Visiting the Web App in Your Browser")
    urls = [
        ("Web Frontend (User Interface): ", "http://127.0.0.1:13000"),
        ("FastAPI Backend Documentation: ", "http://127.0.0.1:18080/docs (Interactive Swagger UI)"),
        ("Backend Health Check: ", "http://127.0.0.1:18080/api/v1/health")
    ]
    for b_title, b_desc in urls:
        add_bullet_point(doc, b_desc, bold_prefix=b_title)

    add_heading_2(doc, "5.3 Demo Script for Presentations (Viva / Review)")
    demo_steps = [
        ("Step 1: Sign Up / Sign In: ", "Go to http://127.0.0.1:13000/login and sign in."),
        ("Step 2: Upload Resume: ", "Go to Profile, upload a sample PDF resume. Show how the skills, education, and experience fields fill in automatically."),
        ("Step 3: Show Job Recommendations: ", "Go to Home / Matches. Show the match score (e.g. 88%) and click 'Why this matches' to show the transparent breakdown."),
        ("Step 4: Demonstrate Auto-Fill: ", "Select a job and click 'Prepare Application'. Show how your name, email, and resume are pre-filled."),
        ("Step 5: Show 'Answer Once': ", "Fill in a missing question (like LinkedIn URL or Work Auth), check 'Save for reuse', and show that future applications remember it!"),
        ("Step 6: Show Tracking: ", "Go to Applications to show the application tracker with statuses and timestamps.")
    ]
    for b_title, b_desc in demo_steps:
        add_bullet_point(doc, b_desc, bold_prefix=b_title)

    add_heading_2(doc, "5.4 Proof of Testing & Code Quality")
    add_body_paragraph(
        doc,
        "If an examiner asks if the code is tested, you can confidently share our verified test counts:"
    )
    tests = [
        ("385 Backend Unit Tests Passing: ", "Run 'python -m pytest backend/tests -q' (covers auth, profiles, jobs, applications)."),
        ("64 PostgreSQL Concurrency Tests: ", "Verifies multi-user database transactions."),
        ("40 Browser Automation Tests: ", "Verifies Playwright form filling and request guarding."),
        ("33 End-to-End Playwright Audit Checks: ", "All documented in BUG_AUDIT.md with zero unresolved critical bugs.")
    ]
    for b_title, b_desc in tests:
        add_bullet_point(doc, b_desc, bold_prefix=b_title)

    # Save Document
    output_path = r"D:\College\4-2 project\TURIO_PROJECT_DOCUMENTATION.docx"
    print(f"Saving easy-to-understand documentation to {output_path}...")
    doc.save(output_path)
    print("Team-friendly document saved successfully.")


if __name__ == "__main__":
    build_team_documentation()
